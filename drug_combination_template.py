"""
Drug Combination Excel Template Generator

Generates an Excel file with worksheets for drug combination experiments.
Each worksheet contains a template for entering drug combination results.

Input: Excel file with drug concentrations and experiment names
Output: Excel file with one worksheet per experiment (all drug pairs combined)
"""

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from itertools import combinations


def read_input_excel(input_path: str):
    """
    Read drug concentrations and experiments from input Excel file.
    
    Args:
        input_path: Path to input Excel file
        
    Returns:
        tuple: (drugs_dict, experiments_list)
            drugs_dict: {drug_name: [conc1, conc2, ...]}
            experiments_list: [{"name": exp_name, "replicates": N}, ...]
    """
    wb = load_workbook(input_path)
    
    # Read DrugAndConcentrations sheet
    drugs = {}
    if "DrugAndConcentrations" in wb.sheetnames:
        ws = wb["DrugAndConcentrations"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            drug_name = row[0]
            if drug_name:
                concentrations = [c for c in row[1:] if c is not None]
                if concentrations:
                    drugs[drug_name] = sorted(concentrations)
    
    # Read ExperimentName sheet
    experiments = []
    if "ExperimentName" in wb.sheetnames:
        ws = wb["ExperimentName"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            exp_name = row[0]
            replicates = row[1] if row[1] else 3
            if exp_name:
                experiments.append({
                    "name": exp_name,
                    "replicates": int(replicates)
                })
    
    wb.close()
    return drugs, experiments


def create_experiment_worksheet(
    wb: Workbook,
    exp: dict,
    drugs: dict,
    drug_pairs: list
):
    """
    Create a single worksheet for one experiment containing all drug pairs.
    
    Structure:
        - Row 1: Header (Single drugs | Drugs | Rep1 | Rep2 | ...)
        - Row 2: Control (Control | 0 | empty | empty | ...)
        - Single drugs section (Drug name | concentration | empty cells)
        - Combination section (Drug_Conc | Drug_Conc | empty cells)
    """
    sheet_name = exp["name"]
    # Truncate if too long (Excel limit is 31 chars)
    if len(sheet_name) > 31:
        sheet_name = sheet_name[:31]
    
    ws = wb.create_sheet(title=sheet_name)
    replicates = exp["replicates"]
    
    # Style definitions
    header_font = Font(bold=True)
    section_font = Font(bold=True)
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    control_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    combination_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    center_align = Alignment(horizontal='center', vertical='center')
    
    current_row = 1
    
    # Header row
    headers = ["Single drugs", "Drugs"] + [f"Rep{i+1}" for i in range(replicates)]
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=current_row, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = center_align
    
    current_row += 1
    
    # Control row
    cell_a = ws.cell(row=current_row, column=1, value="Control")
    cell_a.font = section_font
    cell_a.fill = control_fill
    cell_a.border = thin_border
    cell_a.alignment = center_align
    
    cell_b = ws.cell(row=current_row, column=2, value=0)
    cell_b.fill = control_fill
    cell_b.border = thin_border
    cell_b.alignment = center_align
    
    for col_idx in range(3, 3 + replicates):
        cell = ws.cell(row=current_row, column=col_idx, value="")
        cell.border = thin_border
    
    current_row += 1
    
    # Single drugs section
    for drug_name in drugs.keys():
        concentrations = drugs[drug_name]
        first_conc = True
        
        for conc in concentrations:
            if first_conc:
                cell_a = ws.cell(row=current_row, column=1, value=drug_name)
                cell_a.font = section_font
                cell_a.border = thin_border
                cell_a.alignment = center_align
                first_conc = False
            else:
                cell_a = ws.cell(row=current_row, column=1, value="")
                cell_a.border = thin_border
            
            cell_b = ws.cell(row=current_row, column=2, value=conc)
            cell_b.border = thin_border
            cell_b.alignment = center_align
            
            for col_idx in range(3, 3 + replicates):
                cell = ws.cell(row=current_row, column=col_idx, value="")
                cell.border = thin_border
            
            current_row += 1
    
    # Combination section
    for drug_a, drug_b in drug_pairs:
        conc_a_list = drugs[drug_a]
        conc_b_list = drugs[drug_b]
        
        first_combo_of_pair = True
        
        for conc_a in conc_a_list:
            first_conc_a = True
            
            for conc_b in conc_b_list:
                if first_combo_of_pair:
                    cell_a = ws.cell(row=current_row, column=1, value=f"{drug_a}_{conc_a}")
                    cell_a.fill = combination_fill
                    cell_a.border = thin_border
                    cell_a.alignment = center_align
                    first_combo_of_pair = False
                    first_conc_a = False
                elif first_conc_a:
                    cell_a = ws.cell(row=current_row, column=1, value=f"{drug_a}_{conc_a}")
                    cell_a.fill = combination_fill
                    cell_a.border = thin_border
                    cell_a.alignment = center_align
                    first_conc_a = False
                else:
                    cell_a = ws.cell(row=current_row, column=1, value="")
                    cell_a.fill = combination_fill
                    cell_a.border = thin_border
                
                cell_b = ws.cell(row=current_row, column=2, value=f"{drug_b}_{conc_b}")
                cell_b.fill = combination_fill
                cell_b.border = thin_border
                cell_b.alignment = center_align
                
                for col_idx in range(3, 3 + replicates):
                    cell = ws.cell(row=current_row, column=col_idx, value="")
                    cell.border = thin_border
                
                current_row += 1
    
    # Adjust column widths
    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 15
    for col_idx in range(3, 3 + replicates):
        ws.column_dimensions[get_column_letter(col_idx)].width = 12
    
    return ws


def generate_from_excel(
    input_path: str,
    output_path: str = "drug_combination_results.xlsx"
):
    """
    Read drug concentrations and experiments from Excel,
    generate output file with one worksheet per experiment.
    
    Args:
        input_path: Path to input Excel file (GenExcel.xlsx format)
        output_path: Output Excel file path
        
    Input Excel format:
        Sheet 1: DrugAndConcentrations
            | NAME  | Conc1 | Conc2 | Conc3 | ...
            | DrugA | 0.78  | 1.56  | 3.12  | ...
            | DrugB | 0.09  | 0.19  | 0.39  | ...
            
        Sheet 2: ExperimentName
            | Experiment name | Replicates |
            | Experiment 1    | 4          |
            | Experiment 2    | 3          |
    
    Output:
        One worksheet per experiment containing ALL drug pair combinations
    
    Structure per worksheet:
        1. Header row
        2. Control row (Control | 0)
        3. Single drugs section (each drug with all concentrations)
        4. Combination section (all unique drug pairs)
    """
    # Read input
    drugs, experiments = read_input_excel(input_path)
    
    if not drugs:
        raise ValueError("No drugs found in input file")
    if not experiments:
        raise ValueError("No experiments found in input file")
    
    # Generate all unique drug pairs
    drug_names = list(drugs.keys())
    drug_pairs = list(combinations(drug_names, 2))
    
    print(f"Found {len(drugs)} drugs: {list(drugs.keys())}")
    print(f"Found {len(experiments)} experiments: {[e['name'] for e in experiments]}")
    print(f"Generating {len(drug_pairs)} drug pairs: {drug_pairs}")
    
    # Create workbook
    wb = Workbook()
    wb.remove(wb.active)
    
    # Create one worksheet per experiment
    for exp in experiments:
        create_experiment_worksheet(wb, exp, drugs, drug_pairs)
        print(f"Created worksheet: {exp['name']}")
    
    # Save workbook
    wb.save(output_path)
    print(f"\nExcel file saved to: {output_path}")
    print(f"Total worksheets: {len(experiments)}")
    
    return output_path


if __name__ == "__main__":
    import sys
    
    if len(sys.argv) > 1:
        input_file = sys.argv[1]
        output_file = sys.argv[2] if len(sys.argv) > 2 else "drug_combination_results.xlsx"
    else:
        input_file = "GenExcel.xlsx"
        output_file = "drug_combination_results.xlsx"
    
    generate_from_excel(input_file, output_file)
